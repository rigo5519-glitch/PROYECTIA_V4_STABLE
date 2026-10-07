from docx import Document


class PrefeasibilityGenerator:

    def generate(
        self,
        project_master,
        output_file="prefactibilidad.docx"
    ):

        doc = Document()

        doc.add_heading(
            "ESTUDIO DE PREFACTIBILIDAD",
            level=1
        )

        doc.add_heading(
            "Resumen Ejecutivo",
            level=2
        )

        doc.add_paragraph(
            project_master["idea"]
            ["idea_original"]
        )

        doc.add_heading(
            "Diagnóstico",
            level=2
        )

        doc.add_paragraph(
            str(
                project_master[
                    "diagnostic"
                ]
            )
        )

        doc.add_heading(
            "Mercado",
            level=2
        )

        doc.add_paragraph(
            str(
                project_master[
                    "market"
                ]
            )
        )

        doc.add_heading(
            "Localización",
            level=2
        )

        doc.add_paragraph(
            str(
                project_master[
                    "localization"
                ]
            )
        )

        doc.add_heading(
            "Inversión",
            level=2
        )

        doc.add_paragraph(
            str(
                project_master[
                    "investment"
                ]
            )
        )

        doc.add_heading(
            "VAN",
            level=2
        )

        doc.add_paragraph(
            str(
                project_master[
                    "van"
                ]
            )
        )

        doc.save(
            output_file
        )

        return output_file