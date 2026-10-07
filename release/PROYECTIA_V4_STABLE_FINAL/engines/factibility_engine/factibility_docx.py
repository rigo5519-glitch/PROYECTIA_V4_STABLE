from docx import Document


class FactibilityDOCX:

    def generate(
        self,
        factibilidad,
        output_file
    ):

        doc = Document()

        doc.add_heading(
            "ESTUDIO DE FACTIBILIDAD",
            level=1
        )

        resumen = factibilidad[
            "resumen_ejecutivo"
        ]

        doc.add_heading(
            "1. Resumen Ejecutivo",
            level=2
        )

        doc.add_paragraph(
            f"Idea: {resumen['idea']}"
        )

        doc.add_paragraph(
            f"Problema: {resumen['problema']}"
        )

        doc.add_paragraph(
            f"Objetivo: {resumen['objetivo']}"
        )

        doc.add_heading(
            "2. Mercado",
            level=2
        )

        doc.add_paragraph(
            str(
                factibilidad["mercado"]
            )
        )

        doc.add_heading(
            "3. Marketing",
            level=2
        )

        doc.add_paragraph(
            str(
                factibilidad["marketing"]
            )
        )

        doc.add_heading(
            "4. Estudio Técnico",
            level=2
        )

        doc.add_paragraph(
            str(
                factibilidad["tecnico"]
            )
        )

        financiero = factibilidad[
            "financiero"
        ]

        doc.add_heading(
            "5. Evaluación Financiera",
            level=2
        )

        doc.add_paragraph(
            f"VAN: {financiero['van']['VAN']}"
        )

        doc.add_paragraph(
            f"TIR: {financiero['tir']['TIR']}"
        )

        doc.add_paragraph(
            f"PRI: {financiero['pri']}"
        )

        doc.add_paragraph(
            f"BC: {financiero['bc']['BC']}"
        )

        doc.add_heading(
            "6. Riesgos",
            level=2
        )

        doc.add_paragraph(
            str(
                factibilidad["riesgos"]
            )
        )

        doc.add_heading(
            "7. Conclusión",
            level=2
        )

        doc.add_paragraph(
            factibilidad[
                "conclusion"
            ]
        )

        doc.save(
            output_file
        )