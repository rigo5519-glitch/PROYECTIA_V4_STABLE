from docx import Document

from engines.copilot_engine.copilot_engine import (
    CopilotEngine
)


class CopilotReportGeneratorV2:

    def generate(
        self,
        idea,
        output_file="copilot_report.docx"
    ):

        copilot = CopilotEngine()

        resultado = copilot.generate_report(
            idea
        )

        doc = Document()

        doc.add_heading(
            "PROYECTIA COPILOT REPORT",
            level=1
        )

        doc.add_heading(
            "Idea del Proyecto",
            level=2
        )

        doc.add_paragraph(
            resultado["idea"]
        )

        doc.add_heading(
            "Proyectos Analizados",
            level=2
        )

        doc.add_paragraph(
            str(
                resultado[
                    "proyectos_analizados"
                ]
            )
        )

        doc.add_heading(
            "Ubicación Recomendada",
            level=2
        )

        doc.add_paragraph(
            resultado[
                "ubicacion_frecuente"
            ]
        )

        doc.add_heading(
            "Inversión Promedio",
            level=2
        )

        doc.add_paragraph(
            f'USD {resultado["inversion_promedio"]:,.2f}'
        )

        doc.add_heading(
            "VAN Promedio",
            level=2
        )

        doc.add_paragraph(
            f'USD {resultado["van_promedio"]:,.2f}'
        )

        doc.add_heading(
            "Factibilidad",
            level=2
        )

        doc.add_paragraph(
            resultado["factibilidad"]
        )

        doc.add_heading(
            "Recomendación",
            level=2
        )

        doc.add_paragraph(
            resultado["recomendacion"]
        )

        doc.save(
            output_file
        )

        return output_file