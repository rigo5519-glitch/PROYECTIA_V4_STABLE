from docx import Document

from engines.copilot_engine.executive_brief_generator import (
    ExecutiveBriefGenerator
)


class ExecutiveBriefDOCX:

    def generate(
        self,
        idea,
        output_file="executive_brief.docx"
    ):

        generator = (
            ExecutiveBriefGenerator()
        )

        brief = generator.generate(
            idea
        )

        doc = Document()

        doc.add_heading(
            "PROYECTIA EXECUTIVE BRIEF",
            level=1
        )

        doc.add_paragraph(
            brief
        )

        doc.save(
    output_file
    )

        print(
    "DOCX GUARDADO:",
    output_file
)

        return output_file


    def generate(
        self,
        idea,
        output_file="executive_brief.docx"
    ):

        print("PASO 1")

        generator = ExecutiveBriefGenerator()

        print("PASO 2")

        brief = generator.generate(
            idea
        )

        print("PASO 3")

        doc = Document()

        doc.add_heading(
            "PROYECTIA EXECUTIVE BRIEF",
            level=1
        )

        doc.add_paragraph(
            brief
        )

        print("PASO 4")

        doc.save(
            output_file
        )

        print("PASO 5")

        return output_file